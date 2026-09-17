import os
import re
import shutil
import io
from datetime import datetime
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from PIL import Image
Image.MAX_IMAGE_PIXELS = None
from PyQt6.QtCore import Qt, QAbstractTableModel, QThread, pyqtSignal
from src.config import HEADERS



class DataLoaderThread(QThread):
    progress = pyqtSignal(int, int, str)
    finished = pyqtSignal(list, list)

    def __init__(self, folder_path, scan_all, start_dt, end_dt):
        super().__init__()
        self.folder_path = folder_path
        self.scan_all = scan_all
        self.start_dt = start_dt
        self.end_dt = end_dt

    def load_review_data(self, r_file_path):
        review_features = {}
        if not os.path.exists(r_file_path): 
            return review_features
        current_loc_id1 = ""
        try:
            for event, elem in ET.iterparse(r_file_path, events=("start", "end")):
                tag = elem.tag.split('}')[-1]
                if event == "start":
                    if tag == "Location":
                        current_loc_id1 = elem.attrib.get("Identifier1", "")
                    elif tag == "Feature":
                        feat_id = elem.attrib.get("Identifier", "")
                        err = elem.attrib.get("ReworkError", "")
                        if current_loc_id1 and err: 
                            review_features[(current_loc_id1, feat_id)] = err
                elif event == "end":
                    elem.clear() 
        except Exception: pass
        return review_features

    def extract_datetime_from_filename(self, filename):
        match = re.search(r"(\d{2}-\d{2}-\d{4})\.(\d{2})\.(\d{2})\.(\d{2})", filename)
        if match:
            date_str = match.group(1)
            h, m, s = match.group(2), match.group(3), match.group(4)
            dt_str = f"{date_str} {h}:{m}:{s}"
            try: return datetime.strptime(dt_str, "%m-%d-%Y %H:%M:%S")
            except ValueError: return None
        return None

    def run(self):
        raw_data = []
        unique_values = [set() for _ in HEADERS] 
        xml_files = []

        try:
            with os.scandir(self.folder_path) as entries:
                for entry in entries:
                    if entry.is_file() and entry.name.lower().endswith(".xml") and not entry.name.lower().endswith(".r.xml"):
                        if not self.scan_all:
                            file_dt = self.extract_datetime_from_filename(entry.name)
                            if file_dt:
                                if not (self.start_dt <= file_dt <= self.end_dt): continue 
                            else: continue
                        xml_files.append(entry.path)
        except Exception:
            self.finished.emit([], [list(s) for s in unique_values]) 
            return

        total_files = len(xml_files)
        if total_files == 0:
            self.finished.emit([], [list(s) for s in unique_values]) 
            return
        
        for index, file_path in enumerate(xml_files):
            self.progress.emit(index + 1, total_files, os.path.basename(file_path))
            
            r_file_path = file_path.replace(".I.xml", ".R.xml").replace(".i.xml", ".r.xml")
            review_features = None 

            current_user, current_test_date, current_product = "", "", ""
            current_lot, current_board, current_location = "", "", ""
            current_part_number, current_identifier3 = "", ""
            current_loc_insp_info, current_loc_image1 = "", ""

            try:
                for event, elem in ET.iterparse(file_path, events=("start", "end")):
                    tag_name = elem.tag.split('}')[-1]

                    if event == "start":
                        if tag_name == "Panel":
                            current_test_date = elem.attrib.get("EndTime", "")
                            current_lot = elem.attrib.get("LotCode", "")
                            if review_features is None:
                                review_features = self.load_review_data(r_file_path)

                        elif tag_name == "Image":
                            current_board = elem.attrib.get("Name", "").replace("Board ", "Board")
                            
                        elif tag_name == "Location":
                            current_location = elem.attrib.get("Identifier1", "")
                            current_part_number = elem.attrib.get("Identifier2", "").replace("Component\\", "")
                            current_identifier3 = elem.attrib.get("Identifier3", "")  
                            current_loc_insp_info = elem.attrib.get("LocationInspectionInfo1", "")
                            current_loc_image1 = elem.attrib.get("Image1", "")

                    elif event == "end":
                        if tag_name == "UserName": current_user = elem.text
                        elif tag_name == "Assembly": current_product = elem.text
                        elif tag_name == "Feature":
                            feature_status = elem.attrib.get("FeatureStatus", "")
                            if feature_status == "Pass":
                                elem.clear()
                                continue

                            feature_id = elem.attrib.get("Identifier", "")
                            feature_img1 = elem.attrib.get("Image1", "")

                            inspection = elem.attrib.get("FeatureInspectionInfo", "Not Found")
                            
                            x_val, y_val = "", ""
                            feature_result = elem.find(".//{*}FeatureResult/{*}Measurements")
                            if feature_result is not None:
                                x_elem, y_elem = feature_result.find("{*}X"), feature_result.find("{*}Y")
                                if x_elem is not None: x_val = x_elem.attrib.get("Value", "0.000")
                                if y_elem is not None: y_val = y_elem.attrib.get("Value", "0.000")

                            review_status = "Accept"
                            key = (current_location, feature_id)
                            
                            if review_features is not None and key in review_features:
                                err = review_features[key]
                                if err and err != "False Fail":
                                    review_status = err

                            dt_obj = datetime.min
                            if current_test_date:
                                try: dt_obj = datetime.strptime(current_test_date, "%m/%d/%Y %H:%M:%S")
                                except ValueError: pass

                            row_tuple = (
                                current_user, current_test_date, current_product, current_board, 
                                current_location, current_loc_insp_info, 
                                current_loc_image1, feature_img1,
                                feature_id, current_part_number, current_identifier3, 
                                current_lot, inspection, review_status, current_user, 
                                f"{float(x_val):.3f}" if x_val else "", 
                                f"{float(y_val):.3f}" if y_val else "",
                                dt_obj
                            )
                            raw_data.append(row_tuple)

                            for i, val in enumerate(row_tuple[:-1]):
                                if val: unique_values[i].add(val)
                            
                            elem.clear()
            except Exception: pass
            
        self.finished.emit(raw_data, [sorted(list(s)) for s in unique_values])

class ExcelExportThread(QThread):
    progress = pyqtSignal(int, int, str)
    finished = pyqtSignal(bool, str)

    def __init__(self, export_data, save_path, image_folder, selected_cols):
        super().__init__()
        self.export_data = export_data
        self.save_path = save_path
        self.image_folder = image_folder
        self.selected_cols = selected_cols 

    def process_image(self, img_path, max_dim=800, quality=70):
        try:
            with Image.open(img_path) as pil_img:
                if pil_img.mode != "RGB":
                    pil_img = pil_img.convert("RGB")
                if pil_img.width > max_dim or pil_img.height > max_dim:
                    pil_img.thumbnail((max_dim, max_dim), Image.Resampling.LANCZOS)
                buf = io.BytesIO()
                pil_img.save(buf, format="JPEG", quality=quality, optimize=True)
                buf.seek(0)
                return buf
        except Exception:
            return None

    def run(self):
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "AOI Export"

            img1_col_idx_in_model = HEADERS.index("Review Image")
            imgR_col_idx_in_model = HEADERS.index("Target Image")

            has_image_col = False
            for excel_col_idx, model_col_idx in enumerate(self.selected_cols, 1):
                cell = ws.cell(row=1, column=excel_col_idx, value=HEADERS[model_col_idx])
                cell.font = openpyxl.styles.Font(bold=True)
                if model_col_idx in [img1_col_idx_in_model, imgR_col_idx_in_model]:
                    ws.column_dimensions[get_column_letter(excel_col_idx)].width = 20
                    has_image_col = True

            total_rows = len(self.export_data)

            for row_idx, row_data in enumerate(self.export_data, 2):
                self.progress.emit(row_idx - 1, total_rows, "Exporting Excel...")
                if has_image_col:
                    ws.row_dimensions[row_idx].height = 80

                for excel_col_idx, model_col_idx in enumerate(self.selected_cols, 1):
                    val = row_data[model_col_idx]
                    
                    if model_col_idx in [img1_col_idx_in_model, imgR_col_idx_in_model]:
                        if val:
                            file_name = val if val.lower().endswith(".jpg") else f"{val}.jpg"
                            img_path = os.path.join(self.image_folder, file_name)
                            if os.path.exists(img_path):
                                img_buf = self.process_image(img_path, max_dim=800, quality=70)
                                if img_buf:
                                    img = XLImage(img_buf)
                                    img.width, img.height = 100, 100
                                    ws.add_image(img, f"{get_column_letter(excel_col_idx)}{row_idx}")
                                else:
                                    ws.cell(row=row_idx, column=excel_col_idx, value="Image Error")
                            else:
                                ws.cell(row=row_idx, column=excel_col_idx, value="No Image")
                    else:
                        ws.cell(row=row_idx, column=excel_col_idx, value=str(val))

            wb.save(self.save_path)
            self.finished.emit(True, f"Save in:\n{self.save_path}")
        except Exception as e:
            self.finished.emit(False, f"Error: {str(e)}")

class ImageExportThread(QThread):
    progress = pyqtSignal(int, int, str)
    finished = pyqtSignal(bool, str)

    def __init__(self, export_data, dest_folder, image_folder, target_col_idx):
        super().__init__()
        self.export_data = export_data
        self.dest_folder = dest_folder
        self.image_folder = image_folder
        self.target_col_idx = target_col_idx

    def run(self):
        try:
            total_rows = len(self.export_data)
            saved_count = 0
            
            for row_idx, row_data in enumerate(self.export_data, 1):
                self.progress.emit(row_idx, total_rows, "Extract images...")
                img_name = row_data[self.target_col_idx]
                
                if img_name:
                    file_name = img_name if img_name.lower().endswith(".jpg") else f"{img_name}.jpg"
                    src_path = os.path.join(self.image_folder, file_name)
                    dst_path = os.path.join(self.dest_folder, file_name)
                    
                    if os.path.exists(src_path):
                        shutil.copy2(src_path, dst_path)
                        saved_count += 1
                        
            self.finished.emit(True, f"Successfully saved {saved_count} images to:\n{self.dest_folder}")
        except Exception as e:
            self.finished.emit(False, f"Image export error: {str(e)}")

class FastTableModel(QAbstractTableModel):
    def __init__(self):
        super().__init__()
        self._raw_data, self._display_data = [], []

    def load_data(self, data):
        self.beginResetModel()
        self._raw_data = self._display_data = data
        self.endResetModel()

    def filter_data(self, active_filters, date_range=None):
        self.beginResetModel()
        filtered = self._raw_data
        if active_filters:
            filtered = [row for row in filtered if all(row[col_idx] in valid_vals for col_idx, valid_vals in active_filters.items())]
        if date_range:
            start_dt, end_dt = date_range
            filtered = [row for row in filtered if start_dt <= row[-1] <= end_dt]
        self._display_data = filtered
        self.endResetModel()

    def sort(self, column, order):
        self.layoutAboutToBeChanged.emit()
        self._display_data.sort(key=lambda x: x[column], reverse=(order == Qt.SortOrder.DescendingOrder))
        self.layoutChanged.emit()

    def rowCount(self, parent=None): return len(self._display_data)
    def columnCount(self, parent=None): return len(HEADERS)

    def data(self, index, role):
        if role == Qt.ItemDataRole.DisplayRole:
            return self._display_data[index.row()][index.column()]

    def headerData(self, section, orientation, role):
        if role == Qt.ItemDataRole.DisplayRole:
            return HEADERS[section] if orientation == Qt.Orientation.Horizontal else str(section + 1)